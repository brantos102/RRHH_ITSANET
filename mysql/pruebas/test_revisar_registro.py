"""La revisión del registro de vacaciones, sobre un libro inventado.

Cada problema que se encontró en el registro real se siembra aquí en una
persona ficticia, y la revisión tiene que señalarlo.
"""
from __future__ import annotations

import datetime as dt
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import revisar_registro as rr  # noqa: E402

HOY = dt.date(2026, 10, 26)


@pytest.fixture
def libro(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    m = wb.active
    m.title = rr.HOJA_MAESTRO
    m.append(["NOMBRE", "FECHA DE INGRESO", "CIUDAD ", "BODEGA ", "ESTADO"])
    m.append(["PEREZ ANA", dt.datetime(2020, 8, 17), "GYE", "CONECEL", "ACTIVO"])
    m.append(["ROJAS LUIS", dt.datetime(2015, 3, 1), "UIO", "DIFARE", "ACTIVO"])
    m.append(["VEGA SOL", dt.datetime(2019, 1, 10), "UIO", "DIFARE", "ACTIVO"])   # nunca tomó

    r = wb.create_sheet(rr.HOJA_NUEVA)
    r.append(["NOMBRE", "ESTADO", "FECHA INGRESO", "FECHA ACTUAL", "CIUDAD", "BODEGA", "MESES TRABAJADOS",
              "AÑOS COMPLETOS", "DIAS BASE", "DIAS ADICIONALES", "TOTAL", "inicio", "fin", "DIAS", "DISPONIBLE"])

    def fila(nombre, ini, fin, hoy_a_mano=False, dias_a_mano=None):
        n = r.max_row + 1
        r.append([nombre, f"=VLOOKUP(A{n},X,5,0)", f"=VLOOKUP(A{n},X,2,0)",
                  dt.datetime(2026, 7, 13) if hoy_a_mano else "=TODAY()", f"=VLOOKUP(A{n},X,3,0)",
                  f"=VLOOKUP(A{n},X,4,0)", "=1", "=1", "=1", "=1", "=1", ini, fin,
                  dias_a_mano if dias_a_mano is not None else f"=M{n}-L{n}+1", "=1"])

    d = dt.datetime
    fila("PEREZ ANA", d(2025, 1, 6), d(2025, 1, 10))
    fila("PEREZ ANA", d(2025, 1, 6), d(2025, 1, 10))                      # repetida
    fila("PEREZ ANA", d(2025, 3, 2), d(2025, 3, 9))
    fila("PEREZ ANA", d(2025, 3, 8), d(2025, 3, 12))                      # se pisa 2 días
    fila("PEREZ ANA", d(2026, 1, 20), d(2026, 1, 20), hoy_a_mano=True)    # FECHA ACTUAL congelada
    fila("ROJAS LUIS", d(2014, 5, 1), d(2014, 5, 3))                      # antes del ingreso
    fila("ROJAS LUIS", d(2026, 8, 5), d(2026, 8, 12))                     # incluye el 10 de agosto
    fila("ROJAS LUIS", d(2026, 12, 1), d(2026, 12, 15))                   # futura
    fila("ROJAS LUIS", d(2025, 6, 1), d(2025, 6, 3), dias_a_mano=5)       # días que no cuadran
    fila("ROJAS LUISS", d(2025, 9, 1), d(2025, 9, 2))                     # nombre mal escrito

    a = wb.create_sheet(rr.HOJA_ANTIGUA)
    for _ in range(4):
        a.append([None])
    a.append([1, 926514316, "PEREZ ANA", d(2021, 8, 17)])                 # cédula sin el 0; otro ingreso
    a.append([2, 1700000002, "ROJAS LUIS", d(2015, 3, 1)])                # cédula inválida
    ruta = tmp_path / "registro.xlsx"
    wb.save(ruta)
    return ruta


def _revision(ruta):
    f, v = rr.leer(ruta)
    personas = rr.maestro(v)
    obs = rr.revisar(personas, rr.registro(f, v), rr.antigua(v), HOY)
    return personas, obs, rr.antigua(v)


def _hay(obs, texto, fila=None):
    return any(texto in o[4] and (fila is None or o[2] == fila) for o in obs)


def test_senala_cada_problema_sembrado(libro):
    _, obs, _ = _revision(libro)
    assert _hay(obs, "REPETIDA", 3)
    assert _hay(obs, "Se PISA con la fila 4: 2 día(s)", 5)
    assert _hay(obs, "D FECHA ACTUAL está escrita a mano", 6)
    assert _hay(obs, "anterior a la fecha de ingreso", 7)
    assert _hay(obs, "Incluye feriado", 8)
    assert _hay(obs, "FUTURA", 9)
    assert _hay(obs, "no coinciden con las fechas", 10)
    assert _hay(obs, "Días escritos a mano", 10)
    assert _hay(obs, "no está en FECHA DE INGRESO", 11)
    assert _hay(obs, "ninguna vacación registrada")
    assert _hay(obs, "Fecha de ingreso distinta")
    assert _hay(obs, "Cédula no válida (1700000002)")


def test_la_cedula_recupera_el_cero_inicial(libro):
    _, _, bloques = _revision(libro)
    assert bloques[0]["cedula"] == "0926514316" and rr.cedula_valida("0926514316")


def test_lo_grave_va_primero(libro):
    _, obs, _ = _revision(libro)
    gravedades = [rr.GRAVEDAD[o[0]] for o in obs]
    assert gravedades == sorted(gravedades)


def test_el_libro_de_revision_se_genera(libro):
    personas, obs, bloques = _revision(libro)
    f, v = rr.leer(libro)
    destino = libro.with_name("revision.xlsx")
    rr.escribir(destino, personas, rr.registro(f, v), bloques, obs, HOY)
    wb = openpyxl.load_workbook(destino)
    assert wb.sheetnames == ["Observaciones", "Saldos", "Hoja antigua"]
    assert wb["Saldos"]["B2"].value.date() == HOY
    assert wb["Saldos"]["J5"].value.startswith("=IF(")
