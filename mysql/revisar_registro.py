"""Revisa el registro de vacaciones de Talento Humano antes de cargarlo.

    python mysql/revisar_registro.py REGISTRO_DE_VACACIONES.xlsx

Genera, junto al archivo, «<nombre>_REVISION.xlsx» con tres hojas:

  Observaciones   Cada problema encontrado: hoja, fila, persona, qué pasa,
                  cuántos días mueve y qué hacer. Ordenado por gravedad.
  Saldos          Una fila por persona con FÓRMULAS: lo devengado según la
                  regla de la hoja y según la del sistema, a la fecha de
                  corte que se escriba arriba (celda amarilla). Si las dos
                  reglas no dan lo mismo, la diferencia queda a la vista.
  Hoja antigua    La pestaña «registro de vacaciones» (por períodos), que
                  trae la CÉDULA de cada persona: con ella se empareja al
                  cargar, y aquí se ve si su fecha de ingreso coincide con
                  la de «FECHA DE INGRESO».

No modifica el archivo original ni escribe en ninguna base. El resultado
lleva nombres y cédulas reales: se queda en el equipo de quien lo genera,
no se sube al repositorio ni se comparte fuera de Talento Humano.

LO QUE HAY QUE SABER DE LA HOJA (revisado sobre la versión de octubre 2026)

  «REGISTRO DE VACACIONE» calcula el saldo de cada fila así:
      devengado  = meses completos desde el ingreso × 1,25
                 + días adicionales por antigüedad
      disponible = devengado − suma de los días gozados hasta esa fila
  y toma la fecha de ingreso de «FECHA DE INGRESO» buscando por NOMBRE.

  El sistema calcula exactamente lo mismo (15 días por año, devengados mes
  a mes; desde el año 6, un día más por año, hasta 15 adicionales), con
  una diferencia: la fórmula de la hoja deja de sumar días adicionales a
  partir de los 20 años de servicio (MIN(años; 20)), cuando la ley los
  sigue dando cada año. Hoy nadie supera los 20 años; esta revisión lo
  vigila.

  Los días gozados de la hoja son días CALENDARIO (fin − inicio + 1): un
  feriado dentro de las vacaciones cuenta como gozado. El sistema, para las
  solicitudes nuevas, no lo descuenta. Lo que se carga desde la hoja se
  carga con los días que la hoja dice, para que el saldo cuadre.
"""
from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict
from pathlib import Path

HOJA_NUEVA = "REGISTRO DE VACACIONE"
HOJA_ANTIGUA = "registro de vacaciones"
HOJA_MAESTRO = "FECHA DE INGRESO"

# Feriados nacionales de Ecuador que el sistema tiene cargados (2026), los
# mismos de mysql/migraciones/0007_feriados.sql. Sirven
# para señalar las vacaciones que los incluyen, no para cambiar ningún saldo.
FERIADOS = {dt.date(2026, 1, 1), dt.date(2026, 2, 16), dt.date(2026, 2, 17), dt.date(2026, 4, 3),
            dt.date(2026, 5, 1), dt.date(2026, 5, 24), dt.date(2026, 8, 10), dt.date(2026, 10, 9),
            dt.date(2026, 11, 2), dt.date(2026, 11, 3), dt.date(2026, 12, 25)}

GRAVEDAD = {"ALTA": 1, "MEDIA": 2, "BAJA": 3}

# Columnas de «REGISTRO DE VACACIONE» que tienen que ser fórmula en todas las
# filas. Un valor escrito encima de una fórmula no da error: congela ese
# dato. Así se encontró una «FECHA ACTUAL» escrita a mano (13/07/2026) que
# dejaba el saldo de una persona calculado a julio, 5 días por debajo.
COLUMNAS_FORMULA = {1: "B ESTADO", 2: "C FECHA INGRESO", 3: "D FECHA ACTUAL", 4: "E CIUDAD",
                    5: "F BODEGA", 6: "G MESES TRABAJADOS", 7: "H AÑOS COMPLETOS",
                    8: "I DIAS BASE", 9: "J DIAS ADICIONALES", 10: "K TOTAL DE DIAS",
                    14: "O DISPONIBLE"}


def clave(texto) -> str:
    """Nombre comparable: sin tildes, sin espacios de más, en mayúsculas."""
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def fecha(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def meses_completos(desde: dt.date, hasta: dt.date) -> int:
    """SIFECHA(desde; hasta; "M")."""
    m = (hasta.year - desde.year) * 12 + hasta.month - desde.month
    if hasta.day < desde.day:
        m -= 1
    return max(m, 0)


def cedula_valida(c: str) -> bool:
    if not re.fullmatch(r"\d{10}", c):
        return False
    provincia, tercero = int(c[:2]), int(c[2])
    if not (1 <= provincia <= 24 or provincia == 30) or tercero > 5:
        return False
    suma = 0
    for i, d in enumerate(c[:9]):
        p = int(d) * (2 if i % 2 == 0 else 1)
        suma += p - 9 if p > 9 else p
    return (10 - suma % 10) % 10 == int(c[9])


# ------------------------------------------------------------------ lectura
def leer(archivo: Path):
    import openpyxl
    with warnings.catch_warnings():
        # El archivo trae dos celdas con una «fecha» imposible; openpyxl avisa
        # y las trata como error. Se reportan abajo, no hace falta el aviso.
        warnings.simplefilter("ignore")
        formulas = openpyxl.load_workbook(archivo, data_only=False)
        valores = openpyxl.load_workbook(archivo, data_only=True)
    for h in (HOJA_NUEVA, HOJA_MAESTRO):
        if h not in valores.sheetnames:
            raise SystemExit(f"El archivo no tiene la hoja «{h}». Tiene: {', '.join(valores.sheetnames)}")
    return formulas, valores


def maestro(valores) -> dict[str, dict]:
    personas = {}
    for i, r in enumerate(valores[HOJA_MAESTRO].iter_rows(min_row=2, values_only=True), start=2):
        if not r[0]:
            continue
        personas.setdefault(clave(r[0]), {
            "fila": i, "nombre": str(r[0]).strip(), "ingreso": fecha(r[1]), "ciudad": r[2],
            "bodega": r[3], "estado": str(r[4]).strip() if r[4] is not None else None,
            "repeticiones": 0})
        personas[clave(r[0])]["repeticiones"] += 1
    return personas


def registro(formulas, valores) -> list[dict]:
    filas = []
    hf, hv = formulas[HOJA_NUEVA], valores[HOJA_NUEVA]
    for i, (v, f) in enumerate(zip(hv.iter_rows(min_row=2, max_col=15, values_only=True),
                                   hf.iter_rows(min_row=2, max_col=15, values_only=True)), start=2):
        if not v[0] and not any(x is not None for x in v[11:15]):
            continue
        inicio, fin = fecha(v[11]), fecha(v[12])
        dias = v[13] if isinstance(v[13], (int, float)) else None
        # Un archivo guardado sin valores calculados (algunas exportaciones)
        # trae la fórmula =M-L+1 sin su resultado: se calcula aquí, igual.
        if dias is None and isinstance(f[13], str) and f[13].startswith("=") and inicio and fin:
            dias = (fin - inicio).days + 1
        filas.append({
            "fila": i, "nombre": str(v[0]).strip() if v[0] else None, "k": clave(v[0]),
            "estado": v[1], "ingreso": fecha(v[2]), "ingreso_crudo": v[2],
            "inicio": inicio, "fin": fin, "dias": dias, "dias_crudo": v[13],
            "disponible": v[14], "f_dias": f[13], "f_extra": f[9],
            "a_mano": [COLUMNAS_FORMULA[i] for i in COLUMNAS_FORMULA
                       if v[0] and not (isinstance(f[i], str) and f[i].startswith("="))],
        })
    return filas


def antigua(valores) -> list[dict]:
    if HOJA_ANTIGUA not in valores.sheetnames:
        return []
    bloques, actual = [], None
    for i, r in enumerate(valores[HOJA_ANTIGUA].iter_rows(min_row=5, max_col=12, values_only=True), start=5):
        if r[1] is not None and r[2]:
            crudo = str(r[1]).split(".")[0].strip()
            ced = "0" + crudo if re.fullmatch(r"\d{9}", crudo) else crudo
            actual = {"fila": i, "cedula": ced, "cedula_crudo": crudo, "nombre": str(r[2]).strip(),
                      "k": clave(r[2]), "ingreso": fecha(r[3]), "fechas_texto": 0}
            bloques.append(actual)
        elif actual is not None and (isinstance(r[5], str) or isinstance(r[6], str)) \
                and re.search(r"\d", str(r[5] or r[6])):
            actual["fechas_texto"] += 1
    return bloques


# ------------------------------------------------------------- revisión
def revisar(personas, filas, bloques, hoy: dt.date) -> list[tuple]:
    obs: list[tuple] = []

    def nota(gravedad, hoja, fila, nombre, problema, dias, que_hacer):
        obs.append((gravedad, hoja, fila, nombre, problema, dias, que_hacer))

    por = defaultdict(list)
    for f in filas:
        if f["nombre"]:
            por[f["k"]].append(f)

    for f in filas:
        n, fila = f["nombre"], f["fila"]
        # La fecha de ingreso manda la de FECHA DE INGRESO; la de la fila es
        # una copia por búsqueda que puede venir sin calcular.
        if f["k"] in personas and personas[f["k"]]["ingreso"]:
            f["ingreso"] = personas[f["k"]]["ingreso"]
        if not n:
            nota("ALTA", HOJA_NUEVA, fila, "", "Fila con fechas o días pero sin nombre", None,
                 "Completar el nombre o borrar la fila")
            continue
        if f["k"] not in personas:
            nota("ALTA", HOJA_NUEVA, fila, n, "El nombre no está en FECHA DE INGRESO: la búsqueda no "
                 "lo encuentra (error N/D) y esos días no se descuentan de nadie", f["dias"],
                 "Escribir el nombre exactamente como en FECHA DE INGRESO")
        ini, fin = f["inicio"], f["fin"]
        if (ini is None) != (fin is None):
            nota("ALTA", HOJA_NUEVA, fila, n, "Tiene solo una de las dos fechas", f["dias"],
                 "Completar inicio y fin")
        if ini and fin:
            calendario = (fin - ini).days + 1
            if fin < ini:
                nota("ALTA", HOJA_NUEVA, fila, n, "La fecha de fin es anterior a la de inicio",
                     f["dias"], "Corregir las fechas")
            elif calendario > 45:
                nota("ALTA", HOJA_NUEVA, fila, n, f"Una sola vacación de {calendario} días",
                     calendario, "Confirmar las fechas: puede ser un error de año o de mes")
            if f["dias"] is not None and abs(f["dias"] - calendario) > 0.01:
                nota("ALTA", HOJA_NUEVA, fila, n, f"Los días ({f['dias']:g}) no coinciden con las "
                     f"fechas ({calendario})", f["dias"] - calendario, "Revisar cuál de los dos es el correcto")
            if f["ingreso"] and ini < f["ingreso"]:
                nota("ALTA", HOJA_NUEVA, fila, n, f"Vacación anterior a la fecha de ingreso "
                     f"({f['ingreso']:%d/%m/%Y})", f["dias"],
                     "Revisar la fecha de ingreso o el año de la vacación")
            if ini > hoy:
                nota("MEDIA", HOJA_NUEVA, fila, n, "Vacación FUTURA que ya se descuenta del saldo",
                     f["dias"], "Confirmar que está aprobada; en el sistema irá como solicitud aprobada")
            en_feriado = [ini + dt.timedelta(x) for x in range(max(calendario, 0))
                          if ini + dt.timedelta(x) in FERIADOS]
            if en_feriado:
                nota("MEDIA", HOJA_NUEVA, fila, n, "Incluye feriado: la hoja lo cuenta como día gozado "
                     f"({', '.join(f'{d:%d/%m}' for d in en_feriado)})", len(en_feriado),
                     "Decidir la regla (ver nota): el sistema no descuenta feriados en las solicitudes nuevas")
        if isinstance(f["ingreso_crudo"], str):
            nota("ALTA", HOJA_NUEVA, fila, n, "La fecha de ingreso no se encuentra (error N/D)",
                 None, "Corregir el nombre para que la búsqueda la encuentre")
        if isinstance(f["dias_crudo"], str):
            nota("ALTA", HOJA_NUEVA, fila, n, "Los días gozados dan error", None, "Revisar la fila")
        if isinstance(f["f_extra"], str) and "$C$" in f["f_extra"]:
            ref = int(re.search(r"\$C\$(\d+)", f["f_extra"]).group(1))
            otra = next((g for g in filas if g["fila"] == ref), None)
            misma = otra is not None and otra["k"] == f["k"]
            nota("BAJA" if misma else "ALTA", HOJA_NUEVA, fila, n,
                 f"Los días adicionales usan la fecha de la fila {ref} (referencia fija $C$)"
                 + ("; es la misma persona, no cambia el resultado" if misma else "; ES OTRA PERSONA"),
                 None, "Cambiar $C$ por C en la columna J")
        for col in f["a_mano"]:
            nota("ALTA", HOJA_NUEVA, fila, n, f"La columna {col} está escrita a mano, no es fórmula: "
                 "ese dato queda congelado y el saldo de la fila puede estar mal", None,
                 "Copiar la fórmula de la fila de arriba en esa celda")
        if isinstance(f["f_dias"], (int, float)):
            nota("BAJA", HOJA_NUEVA, fila, n, "Días escritos a mano en vez de la fórmula =M-L+1",
                 None, "Reponer la fórmula para que no se desajuste si cambian las fechas")

    for k, fs in por.items():
        n = fs[0]["nombre"]
        con = [f for f in fs if f["inicio"] and f["fin"] and f["fin"] >= f["inicio"]]
        vistos = defaultdict(list)
        for f in con:
            vistos[(f["inicio"], f["fin"])].append(f["fila"])
        for (ini, fin), lista in vistos.items():
            if len(lista) > 1:
                nota("ALTA", HOJA_NUEVA, lista[1], n, f"Vacación REPETIDA: las filas {lista} tienen las "
                     f"mismas fechas ({ini:%d/%m/%Y}–{fin:%d/%m/%Y}); se descuenta {len(lista)} veces",
                     (fin - ini).days + 1, "Borrar la fila repetida")
        orden = sorted(con, key=lambda f: (f["inicio"], f["fin"]))
        for a, b in zip(orden, orden[1:]):
            if b["inicio"] <= a["fin"] and (a["inicio"], a["fin"]) != (b["inicio"], b["fin"]):
                dobles = (min(a["fin"], b["fin"]) - b["inicio"]).days + 1
                nota("ALTA", HOJA_NUEVA, b["fila"], n, f"Se PISA con la fila {a['fila']}: "
                     f"{dobles} día(s) descontados dos veces", dobles,
                     "Corregir las fechas de una de las dos")
        if [f["fila"] for f in con] != [f["fila"] for f in orden]:
            nota("BAJA", HOJA_NUEVA, fs[0]["fila"], n, "Las vacaciones no están en orden de fecha: "
                 "el saldo de las filas intermedias engaña (el de la última es correcto)",
                 None, "Ordenar por nombre y fecha de inicio")

    for k, p in personas.items():
        if p["repeticiones"] > 1:
            nota("ALTA", HOJA_MAESTRO, p["fila"], p["nombre"], f"Nombre repetido {p['repeticiones']} veces: "
                 "la búsqueda toma siempre la primera", None, "Dejar una sola fila por persona")
        if p["ingreso"] is None:
            nota("ALTA", HOJA_MAESTRO, p["fila"], p["nombre"], "Sin fecha de ingreso", None, "Completarla")
        if k not in por and p["estado"] == "ACTIVO" and p["ingreso"] and meses_completos(p["ingreso"], hoy) >= 12:
            nota("MEDIA", HOJA_MAESTRO, p["fila"], p["nombre"], "Activo con más de un año y ninguna "
                 "vacación registrada", None, "Confirmar que de verdad no ha tomado ninguna")

    for b in bloques:
        p = personas.get(b["k"])
        if not cedula_valida(b["cedula"]):
            nota("ALTA", HOJA_ANTIGUA, b["fila"], b["nombre"], f"Cédula no válida ({b['cedula_crudo']})",
                 None, "Corregir la cédula")
        if p and b["ingreso"] and p["ingreso"] and b["ingreso"] != p["ingreso"]:
            nota("ALTA", HOJA_ANTIGUA, b["fila"], b["nombre"], f"Fecha de ingreso distinta: aquí "
                 f"{b['ingreso']:%d/%m/%Y}, en FECHA DE INGRESO {p['ingreso']:%d/%m/%Y}", None,
                 "Confirmar cuál es la correcta: de ella sale todo lo devengado")
    return sorted(obs, key=lambda o: (GRAVEDAD[o[0]], o[1], o[2]))


# ------------------------------------------------------------- escritura
def escribir(destino: Path, personas, filas, bloques, obs, hoy: dt.date) -> None:
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    letra = Font(name="Arial", size=10)
    negrita = Font(name="Arial", size=10, bold=True)
    titulo = Font(name="Arial", size=12, bold=True)
    entrada = Font(name="Arial", size=10, color="0000FF")
    amarillo = PatternFill("solid", fgColor="FFFF00")
    cabecera = PatternFill("solid", fgColor="D9E1F2")
    colores = {"ALTA": "F8CBAD", "MEDIA": "FFE699", "BAJA": "E2EFDA"}

    wb = openpyxl.Workbook()

    def encabezar(ws, fila, nombres, anchos):
        for c, (n, a) in enumerate(zip(nombres, anchos), start=1):
            celda = ws.cell(fila, c, n)
            celda.font, celda.fill = negrita, cabecera
            celda.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(c)].width = a
        ws.freeze_panes = ws.cell(fila + 1, 1)

    # ---- Observaciones
    ws = wb.active
    ws.title = "Observaciones"
    ws["A1"] = "Revisión del registro de vacaciones"
    ws["A1"].font = titulo
    resumen = Counter(o[0] for o in obs)
    ws["A2"] = (f"Revisado el {hoy:%d/%m/%Y}. {len(obs)} observaciones: {resumen.get('ALTA', 0)} de gravedad "
                f"ALTA (mueven saldos), {resumen.get('MEDIA', 0)} MEDIA (decisiones), "
                f"{resumen.get('BAJA', 0)} BAJA (orden y forma).")
    ws["A2"].font = letra
    encabezar(ws, 4, ["Gravedad", "Hoja", "Fila", "Persona", "Qué pasa", "Días afectados", "Qué hacer"],
              [10, 22, 7, 38, 70, 10, 55])
    for r, o in enumerate(obs, start=5):
        for c, v in enumerate(o, start=1):
            celda = ws.cell(r, c, v)
            celda.font = letra
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=colores[o[0]])
    ws.auto_filter.ref = f"A4:G{max(5, 4 + len(obs))}"

    # ---- Saldos (con fórmulas)
    ws = wb.create_sheet("Saldos")
    ws["A1"] = "Saldo de cada persona: regla de la hoja y regla del sistema"
    ws["A1"].font = titulo
    ws["A2"] = "Fecha de corte:"
    ws["A2"].font = negrita
    ws["B2"] = hoy
    ws["B2"].number_format = "dd/mm/yyyy"
    ws["B2"].font, ws["B2"].fill = entrada, amarillo
    ws["B2"].comment = Comment("Escriba aquí la fecha a la que quiere ver los saldos (por ejemplo, la del "
                               "corte de la migración). Todo lo demás se recalcula.", "Revisión")
    ws["D2"] = ("Celda amarilla: la única que se edita. Devengado = meses completos × 1,25 + adicionales por "
                "antigüedad (desde el año 6, un día más por año, hasta 15). La hoja original deja de sumar "
                "adicionales después de 20 años; el sistema no.")
    ws["D2"].font = letra
    ws["D2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("D2:N2")
    ws.row_dimensions[2].height = 42
    cols = ["Persona", "Estado", "Ciudad", "Bodega", "Ingreso", "Meses completos", "Años completos",
            "Días gozados (hoja)", "Devengado regla hoja", "Devengado regla sistema", "Disponible regla hoja",
            "Disponible regla sistema", "Diferencia", "Último DISPONIBLE guardado en la hoja"]
    encabezar(ws, 4, cols, [38, 9, 8, 18, 11, 9, 8, 10, 11, 11, 11, 11, 10, 14])
    gozados = defaultdict(float)
    ultimo = {}
    for f in filas:
        if f["nombre"]:
            gozados[f["k"]] += f["dias"] or 0
            ultimo[f["k"]] = f["disponible"]
    r = 5
    for k, p in sorted(personas.items(), key=lambda x: x[1]["nombre"]):
        ws.cell(r, 1, p["nombre"])
        ws.cell(r, 2, p["estado"])
        ws.cell(r, 3, p["ciudad"])
        ws.cell(r, 4, p["bodega"])
        ws.cell(r, 5, p["ingreso"]).number_format = "dd/mm/yyyy"
        ws.cell(r, 6, f'=IF(E{r}="","",DATEDIF(E{r},$B$2,"M"))')
        ws.cell(r, 7, f'=IF(E{r}="","",DATEDIF(E{r},$B$2,"Y"))')
        ws.cell(r, 8, gozados.get(k, 0))
        ws.cell(r, 9, f'=IF(E{r}="","",F{r}*1.25+IF(G{r}<=5,0,(MIN(G{r},20)-5)*(MIN(G{r},20)-4)/2))')
        ws.cell(r, 10, f'=IF(E{r}="","",F{r}*1.25+IF(G{r}<=5,0,IF(G{r}<=20,(G{r}-5)*(G{r}-4)/2,120+15*(G{r}-20))))')
        ws.cell(r, 11, f'=IF(E{r}="","",I{r}-H{r})')
        ws.cell(r, 12, f'=IF(E{r}="","",J{r}-H{r})')
        ws.cell(r, 13, f'=IF(E{r}="","",L{r}-K{r})')
        guardado = ultimo.get(k)
        ws.cell(r, 14, guardado if isinstance(guardado, (int, float)) else None)
        for c in range(1, 15):
            ws.cell(r, c).font = letra
        for c in (6, 7, 8, 9, 10, 11, 12, 13, 14):
            ws.cell(r, c).number_format = '#,##0.00;-#,##0.00;"-"'
        r += 1
    total = r
    ws.cell(total, 1, "TOTAL").font = negrita
    for c in (8, 9, 10, 11, 12, 13):
        col = get_column_letter(c)
        ws.cell(total, c, f"=SUM({col}5:{col}{total - 1})").font = negrita
        ws.cell(total, c).number_format = '#,##0.00;-#,##0.00;"-"'
    ws.auto_filter.ref = f"A4:N{total - 1}"

    # ---- Hoja antigua
    ws = wb.create_sheet("Hoja antigua")
    ws["A1"] = "Pestaña «registro de vacaciones»: cédulas y fechas de ingreso"
    ws["A1"].font = titulo
    ws["A2"] = ("La cédula se recupera de esta pestaña. Las que tenían 9 dígitos perdieron el cero inicial "
                "al guardarse como número; aquí se repone.")
    ws["A2"].font = letra
    encabezar(ws, 4, ["Fila", "Cédula (10 dígitos)", "Cédula válida", "Nombre", "Ingreso aquí",
                      "Ingreso en FECHA DE INGRESO", "¿Coincide?", "Está en FECHA DE INGRESO",
                      "Fechas escritas como texto"], [7, 14, 9, 38, 12, 14, 10, 12, 12])
    for r, b in enumerate(bloques, start=5):
        p = personas.get(b["k"])
        valores = [b["fila"], b["cedula"], "sí" if cedula_valida(b["cedula"]) else "NO", b["nombre"],
                   b["ingreso"], p["ingreso"] if p else None, None, "sí" if p else "NO", b["fechas_texto"]]
        for c, v in enumerate(valores, start=1):
            ws.cell(r, c, v).font = letra
        ws.cell(r, 2).number_format = "@"
        ws.cell(r, 5).number_format = ws.cell(r, 6).number_format = "dd/mm/yyyy"
        ws.cell(r, 7, f'=IF(OR(E{r}="",F{r}=""),"",IF(E{r}=F{r},"sí","NO"))').font = letra
    ws.auto_filter.ref = f"A4:I{max(5, 4 + len(bloques))}"

    wb.save(destino)


def principal() -> int:
    ap = argparse.ArgumentParser(description="Revisa el registro de vacaciones antes de cargarlo.")
    ap.add_argument("archivo", type=Path)
    ap.add_argument("--corte", type=dt.date.fromisoformat, default=dt.date.today(),
                    help="fecha de corte, AAAA-MM-DD (por omisión, hoy)")
    args = ap.parse_args()
    if not args.archivo.exists():
        print(f"No existe {args.archivo}")
        return 2
    formulas, valores = leer(args.archivo)
    personas = maestro(valores)
    filas = registro(formulas, valores)
    bloques = antigua(valores)
    obs = revisar(personas, filas, bloques, args.corte)
    destino = args.archivo.with_name(args.archivo.stem + "_REVISION.xlsx")
    escribir(destino, personas, filas, bloques, obs, args.corte)
    g = Counter(o[0] for o in obs)
    print(f"Personas en FECHA DE INGRESO: {len(personas)} · filas de vacaciones: {len(filas)} · "
          f"personas en la hoja antigua: {len(bloques)}")
    print(f"Observaciones: {len(obs)} (ALTA {g.get('ALTA', 0)}, MEDIA {g.get('MEDIA', 0)}, BAJA {g.get('BAJA', 0)})")
    print(f"Revisión guardada en: {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(principal())
