"""Carga a la base el personal ACTIVO del registro de vacaciones de Talento Humano.

    .venv/bin/python mysql/cargar_registro.py REGISTRO.xlsx               # revisa; no escribe
    .venv/bin/python mysql/cargar_registro.py REGISTRO.xlsx --aplicar     # carga
    .venv/bin/python mysql/cargar_registro.py REGISTRO.xlsx --quitar --aplicar
                                                                          # borra lo cargado

Después, el historial de cada vacación (no toca saldos):

    .venv/bin/python scripts/cargar_historial.py REGISTRO.xlsx --aplicar

QUÉ SE CARGA, Y DE DÓNDE
  personas    las de estado ACTIVO de «REGISTRO DE VACACIONE»
  cédula      de las hojas que la traen («registro de vacaciones», «totales
              vacaciones», «PASIVOS», «CONECEL», «PASIVOS CONTROL»), por nombre
              exacto. Se repone el 0 inicial que Excel se come y se valida el
              dígito verificador.
  ingreso,    de la hoja maestra «FECHA DE INGRESO» (si falta ahí, de la fila
  ciudad,     del registro)
  bodega,
  CECO
  saldo       el DISPONIBLE de la última fila de cada persona, TAL COMO ESTÁ en la
              hoja. Si la hoja se guardó antes del día de la carga, se le suma
              solo lo que el sistema devenga entre esa fecha y hoy, para que
              nadie gane ni pierda días por la diferencia de fechas.

QUÉ NO SE CARGA, Y SE DICE
  · quien no tiene cédula en ninguna hoja, o tiene dos distintas, o comparte
    cédula con otro nombre: no hay forma segura de saber quién es.
  · las personas PASIVO: su tratamiento está pendiente de decisión de Talento
    Humano (ver el informe de novedades).
  · un saldo mayor que lo devengado: no es un saldo, es un dato por conciliar.
  · el correo: la hoja no trae ninguno. Cada persona queda «sin correo» hasta
    que Talento Humano lo registre (mysql/cargar_personas.py o Administración).
    Mientras tanto no puede recibir el código de acceso.

La lista de lo que no entró, con nombres, se escribe en un CSV junto al Excel
(en el servidor), nunca en el repositorio.

VOLVER A CARGAR. Se identifica a cada persona por cédula: si ya existe, se
actualizan ingreso, ciudad, bodega, CECO y saldo, y se respetan su correo, su
rol y su jefe. A quien ya tenga vacaciones aprobadas DENTRO del sistema no se le
toca el saldo (se perdería lo aprobado); se lista para revisarlo.
"""
from __future__ import annotations

import argparse
import asyncio
import collections
import csv
import datetime as dt
import re
import sys
import warnings

# Dos celdas del archivo traen una «fecha» imposible; se reportan como dato
# faltante. El aviso de openpyxl sobra en pantalla.
warnings.filterwarnings("ignore", message="Cell .* is marked as a date")
from decimal import Decimal
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "servidor"))
sys.path.insert(0, str(RAIZ / "mysql"))

import revisar_registro as rr  # noqa: E402
from app.cedula import es_cedula_valida  # noqa: E402
from app.db import a_json, cerrar_pool, ejecutar, lista, nuevo_id, obtener_todos, obtener_uno, transaccion  # noqa: E402
from app.errores import ReglaDeNegocio  # noqa: E402
from app.reglas import periodos  # noqa: E402

VERDE, ROJO, AMARILLO, GRIS, FIN = "\033[92m", "\033[91m", "\033[93m", "\033[90m", "\033[0m"

# Hoja → (fila del encabezado, columna de la cédula, columna del nombre)
HOJAS_CON_CEDULA = {
    "registro de vacaciones": (4, 1, 2),
    "totales vacaciones": (1, 0, 1),
    "PASIVOS ": (1, 1, 2),
    "CONECEL": (1, 0, 1),
    "PASIVOS CONTROL": (1, 0, 1),
}


# La hoja anota la ciudad con el código de la sede. La región —y con ella el
# equipo de Talento Humano que recibe las solicitudes— se deduce del nombre
# de la ciudad (tabla ciudades_region): con el código, todos quedaban en
# sierra y las solicitudes de Guayaquil iban a Quito.
CIUDADES = {"GYE": "Guayaquil", "UIO": "Quito"}


class _Deshacer(Exception):
    pass


def _ciudad(valor) -> str | None:
    texto = _texto(valor)
    return CIUDADES.get(texto.upper(), texto) if texto else None


def _cedula(valor) -> str | None:
    if valor is None:
        return None
    digitos = "".join(c for c in str(valor).split(".")[0] if c.isdigit())
    return digitos.zfill(10) if 9 <= len(digitos) <= 10 else None


def _texto(valor) -> str | None:
    t = " ".join(str(valor).split()) if valor is not None else ""
    return t or None


def leer(archivo: Path) -> tuple[list[dict], list[dict], collections.Counter]:
    """(para cargar, no se cargan, resumen)."""
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    if rr.HOJA_NUEVA not in wb.sheetnames:
        raise SystemExit(f"El archivo no tiene la hoja «{rr.HOJA_NUEVA}».")

    # Cédulas por nombre, de todas las hojas que las traen.
    cedulas: dict[str, set[str]] = collections.defaultdict(set)
    for hoja, (encabezado, col_ced, col_nom) in HOJAS_CON_CEDULA.items():
        if hoja not in wb.sheetnames:
            continue
        for f in wb[hoja].iter_rows(min_row=encabezado + 1, values_only=True):
            if len(f) > max(col_ced, col_nom) and f[col_nom]:
                c = _cedula(f[col_ced])
                if c:
                    cedulas[rr.clave(f[col_nom])].add(c)
    nombres_por_cedula: dict[str, set[str]] = collections.defaultdict(set)
    for k, cs in cedulas.items():
        for c in cs:
            nombres_por_cedula[c].add(k)

    maestro = {}
    if rr.HOJA_MAESTRO in wb.sheetnames:
        for r in wb[rr.HOJA_MAESTRO].iter_rows(min_row=2, values_only=True):
            if r and r[0]:
                maestro.setdefault(rr.clave(r[0]), {
                    "ingreso": rr.fecha(r[1]), "ciudad": _texto(r[2]), "bodega": _texto(r[3]),
                    "ceco": _texto(r[5]) if len(r) > 5 else None})

    # La última fila de cada persona trae su saldo vigente.
    ultima: dict[str, tuple[int, tuple]] = {}
    for i, v in enumerate(wb[rr.HOJA_NUEVA].iter_rows(min_row=2, max_col=15, values_only=True), start=2):
        if v and v[0]:
            ultima[rr.clave(v[0])] = (i, v)

    cargar, fuera, resumen = [], [], collections.Counter()
    for k, (fila, v) in ultima.items():
        nombre = _texto(v[0])
        estado = (_texto(v[1]) or "").upper()
        if estado != "ACTIVO":
            resumen["pasivo (pendiente de decisión)"] += 1
            fuera.append({"fila": fila, "nombre": nombre, "motivo": f"estado {estado or 'vacío'}"})
            continue
        m = maestro.get(k, {})
        ingreso = m.get("ingreso") or rr.fecha(v[2])
        cs = cedulas.get(k, set())
        motivo = None
        if not cs:
            motivo = "sin cédula en ninguna hoja"
        elif len(cs) > 1:
            motivo = f"dos cédulas distintas en las hojas ({', '.join(sorted(cs))})"
        else:
            cedula = next(iter(cs))
            if not es_cedula_valida(cedula):
                motivo = f"cédula {cedula} no válida"
            elif len(nombres_por_cedula[cedula]) > 1:
                motivo = f"la cédula {cedula} aparece con otro nombre"
        if motivo is None and ingreso is None:
            motivo = "sin fecha de ingreso"
        disponible = v[14]
        if motivo is None and not isinstance(disponible, (int, float)):
            motivo = f"DISPONIBLE no es un número ({disponible!r})"
        if motivo:
            resumen[re.sub(r" \(.*\)|\d{10}", "…", motivo).replace("……", "…")] += 1
            fuera.append({"fila": fila, "nombre": nombre, "motivo": motivo})
            continue
        cargar.append({
            "fila": fila, "clave": k, "nombre": nombre, "cedula": cedula, "ingreso": ingreso,
            "ciudad": _ciudad(m.get("ciudad") or v[4]), "bodega": m.get("bodega") or _texto(v[5]),
            "ceco": m.get("ceco"), "disponible": Decimal(str(disponible)).quantize(Decimal("0.01")),
            "corte": rr.fecha(v[3]) or dt.date.today()})
    return cargar, fuera, resumen


async def _devengado_al(ingreso: dt.date, corte: dt.date) -> Decimal:
    """Lo que el sistema asigna desde el ingreso hasta `corte` (años cumplidos + año en curso)."""
    fila = await obtener_uno(
        """SELECT anios_cumplidos_al(%s, %s) AS anios""", (ingreso, corte))
    anios = int(fila["anios"] or 0)
    total = Decimal("0")
    for k in range(1, anios + 1):
        total += Decimal(str((await obtener_uno("SELECT dias_por_antiguedad(%s) AS d", (k,)))["d"]))
    total += Decimal(str((await obtener_uno("SELECT dias_devengados_al(%s, %s, %s) AS d",
                                            (ingreso, anios + 1, corte)))["d"]))
    return total


async def _cargar(personas: list[dict], hoy: dt.date) -> dict:
    cuenta = collections.Counter()
    problemas: list[dict] = []
    for p in personas:
        ya = await obtener_uno("SELECT id FROM users WHERE cedula = %s", (p["cedula"],))
        datos = {"nombre": p["nombre"], "fecha_ingreso": p["ingreso"], "ciudad": p["ciudad"],
                 "bodega": p["bodega"], "centro_costo": p["ceco"]}
        if ya is None:
            user_id = str(nuevo_id())
            await ejecutar(
                """INSERT INTO users (id, cedula, nombre, email, correo_pendiente, rol, fecha_ingreso,
                                      ciudad, bodega, centro_costo)
                   VALUES (%s, %s, %s, NULL, TRUE, 'empleado', %s, %s, %s, %s)""",
                (user_id, p["cedula"], p["nombre"], p["ingreso"], p["ciudad"], p["bodega"], p["ceco"]))
            cuenta["nuevas"] += 1
        else:
            user_id = str(ya["id"])
            # Correo, rol y jefe no vienen en la hoja: se respetan.
            await ejecutar(
                """UPDATE users SET nombre = %s, fecha_ingreso = %s, ciudad = %s, bodega = %s,
                                    centro_costo = %s, activo = TRUE WHERE id = %s""",
                (*datos.values(), user_id))
            cuenta["actualizadas"] += 1
            aprobadas = (await obtener_uno(
                """SELECT COUNT(*) AS n FROM requests
                    WHERE user_id = %s AND tipo = 'vacacion' AND estado IN ('aprobado', 'pendiente_anulacion')""",
                (user_id,)))["n"]
            if aprobadas:
                problemas.append({**p, "motivo": f"tiene {aprobadas} vacación(es) aprobada(s) en el "
                                                 "sistema: su saldo no se reemplaza"})
                cuenta["saldo respetado"] += 1
                continue

        # El saldo de la hoja, llevado al día de hoy con lo que el sistema
        # devenga entre la fecha de la hoja y hoy (normalmente, cero).
        saldo = p["disponible"]
        if p["corte"] < hoy:
            saldo += await _devengado_al(p["ingreso"], hoy) - await _devengado_al(p["ingreso"], p["corte"])
        try:
            await periodos.cargar_saldo_inicial(user_id, saldo, 0)
        except ReglaDeNegocio as exc:
            # Más días que los devengados: se generan los períodos (para que la
            # persona exista con su antigüedad) y se deja para conciliar.
            problemas.append({**p, "motivo": f"saldo de la hoja {p['disponible']} no cuadra: {exc.mensaje}"})
            await periodos.generar_periodos_vacaciones(user_id)
            cuenta["saldo por conciliar"] += 1
            continue
        p["saldo_cargado"] = saldo
        cuenta["con saldo"] += 1
    return {"cuenta": cuenta, "problemas": problemas}


def _escribir(destino: Path, filas: list[dict], columnas: list[str]) -> None:
    with destino.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(columnas)
        for r in filas:
            w.writerow([r.get(c, "") if r.get(c) is not None else "" for c in columnas])
    destino.chmod(0o600)


async def principal(args) -> int:
    archivo = Path(args.archivo).expanduser()
    if not archivo.exists():
        raise SystemExit(f"No existe {archivo}")
    print(f"Leyendo {archivo.name}…")
    personas, fuera, resumen = leer(archivo)
    print(f"  {VERDE}{len(personas)}{FIN} persona(s) activas listas para cargar")
    for motivo, n in resumen.most_common():
        print(f"  {AMARILLO}{n:>4}{FIN} no se cargan: {motivo}")
    hoy = dt.date.today()

    try:
        if args.quitar:
            existentes = [f["cedula"] for f in await obtener_todos(
                "SELECT cedula FROM users WHERE cedula IN %s", (lista([p["cedula"] for p in personas]),))]
            print(f"\nSe BORRARÍAN {len(existentes)} persona(s) de este registro con todo su expediente.")
            if not args.aplicar or not existentes:
                print(f"{AMARILLO}Nada se borró.{FIN}" + (" Repita con --aplicar." if existentes else ""))
                return 0
            async with transaccion():
                ids = lista([str(f["id"]) for f in await obtener_todos(
                    "SELECT id FROM users WHERE cedula IN %s", (lista(existentes),))])
                for columna in ("jefe_id", "salida_registrada_por"):
                    await ejecutar(f"UPDATE users SET {columna} = NULL WHERE {columna} IN %s", (ids,))
                await ejecutar("DELETE FROM request_adjustments WHERE ajustado_por IN %s", (ids,))
                n = await ejecutar("DELETE FROM users WHERE id IN %s", (ids,))
                await ejecutar("INSERT INTO audit_logs (accion, entidad, detalle) VALUES ('registro_quitado', 'users', %s)",
                               (a_json({"personas": n, "archivo": archivo.name}),))
            print(f"{VERDE}Borradas {n} persona(s).{FIN}")
            return 0

        # La carga se ensaya siempre dentro de una transacción. Sin --aplicar
        # se deshace al final: así la vista previa es exactamente lo que pasaría.
        resultado = None
        try:
            async with transaccion():
                resultado = await _cargar(personas, hoy)
                if not args.aplicar:
                    raise _Deshacer
                await ejecutar("INSERT INTO audit_logs (accion, entidad, detalle) VALUES ('registro_cargado', 'users', %s)",
                               (a_json({**resultado["cuenta"], "archivo": archivo.name}),))
        except _Deshacer:
            pass
        c = resultado["cuenta"]
        total = sum(p.get("saldo_cargado", 0) for p in personas)
        print(f"\n== {'CARGADO' if args.aplicar else 'LO QUE SE CARGARÍA'} ==")
        print(f"  nuevas: {c['nuevas']} · actualizadas: {c['actualizadas']} · con saldo de la hoja: "
              f"{c['con saldo']} · saldo total: {total} días")
        if c["saldo por conciliar"] or c["saldo respetado"]:
            print(f"  {AMARILLO}saldo por conciliar: {c['saldo por conciliar']} · saldo respetado por "
                  f"vacaciones ya aprobadas: {c['saldo respetado']}{FIN}")

        informe = archivo.with_name(archivo.stem + "_NO_CARGADOS.csv")
        _escribir(informe, [*fuera, *resultado["problemas"]], ["fila", "nombre", "cedula", "motivo"])
        print(f"\n  Lo que no entró o hay que revisar: {informe}  ({len(fuera) + len(resultado['problemas'])} fila(s))")
        if args.aplicar:
            completar = archivo.with_name("personas_completar.csv")
            filas = await obtener_todos(
                """SELECT cedula, nombre, '' AS email, rol, fecha_ingreso, '' AS jefe_cedula, '' AS saldo,
                          cargo, departamento, ciudad
                     FROM users WHERE cedula IN %s AND email IS NULL ORDER BY nombre""",
                (lista([p["cedula"] for p in personas]),))
            _escribir(completar, filas, ["cedula", "nombre", "email", "rol", "fecha_ingreso", "jefe_cedula",
                                         "saldo", "cargo", "departamento", "ciudad"])
            print(f"  Para registrar correos, jefes y roles: {completar}")
            print(f"  (complete esas columnas y cárguelo con mysql/cargar_personas.py; con «saldo» vacío "
                  f"el saldo cargado se conserva)")
            print(f"\n{VERDE}Listo.{FIN} Siguiente: el historial →  .venv/bin/python scripts/cargar_historial.py "
                  f"{archivo} --aplicar")
        else:
            print(f"\n{AMARILLO}Nada se escribió.{FIN} Repita con --aplicar para cargarlo.")
        return 0
    finally:
        await cerrar_pool()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("archivo")
    parser.add_argument("--aplicar", action="store_true", help="escribir de verdad (si no, solo muestra)")
    parser.add_argument("--quitar", action="store_true", help="borrar a las personas de este registro")
    sys.exit(asyncio.run(principal(parser.parse_args())))
