"""Carga el historial de vacaciones de la hoja de Talento Humano.

Talento Humano lleva desde hace años una hoja con cada vacación tomada.
El sistema no tenía nada de eso: un colaborador entraba, veía su saldo y no
podía responder la pregunta más simple que se le ocurre a cualquiera,
«¿cuándo tomé vacaciones la última vez?».

    python scripts/cargar_historial.py ARCHIVO.xlsx            # dice qué haría
    python scripts/cargar_historial.py ARCHIVO.xlsx --aplicar  # lo hace
    python scripts/cargar_historial.py ARCHIVO.xlsx --informe cotejo.csv

Lo delicado aquí es el emparejamiento. La hoja identifica a la gente por
NOMBRE y el sistema por CÉDULA, así que hay que decidir quién es quién, y
equivocarse significa cargarle a una persona las vacaciones de otra. La regla
es estricta a propósito:

    solo se carga cuando el nombre completo coincide EXACTAMENTE
    —ignorando tildes, mayúsculas y espacios de más— con una y solo una
    persona activa del sistema.

Nada de parecidos, nada de «se parece bastante». Lo que no empareja se lista
para que una persona lo mire, y no se carga. Es preferible que falten datos a
que estén mal atribuidos: lo primero se nota, lo segundo no.

Lo cargado NO toca el saldo. El saldo ya viene de esta misma hoja y es
correcto; lo que faltaba era el detalle de cómo se llegó a él. Sumarlo otra
vez sería descontar dos veces los mismos días.
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import datetime as dt
import sys
import unicodedata
from collections import Counter
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "backend"))

VERDE, ROJO, AMARILLO, GRIS, FIN = "\033[92m", "\033[91m", "\033[93m", "\033[90m", "\033[0m"

# La pestaña que manda, tal como se llama hoy: «REGISTRO DE VACACIONE», sin
# la ese final. Es una errata de quien creó el archivo, pero es el nombre
# real y va primero.
#
# HAY OTRA PESTAÑA QUE SE LE PARECE. El mismo archivo trae también «registro
# de vacaciones» —en minúsculas y con ese—, hoy vacía. Cargar la equivocada
# no daría error: daría cero eventos, o los de otra versión de la hoja, y eso
# no se nota. Por eso el nombre exacto gana siempre, y el parecido solo sirve
# cuando el exacto no está y no hay ninguna duda de cuál es.
HOJA = "REGISTRO DE VACACIONE"
HOJA_PREFIJO = "REGISTRO DE VACACION"

# Las columnas de la hoja, por posición. Si cambian de sitio, esto se rompe
# ruidosamente al comprobar el encabezado, que es lo que se quiere.
COL = {"nombre": 0, "estado": 1, "ingreso": 2, "inicio": 11, "fin": 12,
       "dias": 13, "disponible": 14}
ENCABEZADOS = {0: "NOMBRE", 1: "ESTADO", 11: "inicio", 12: "fin"}

# Dónde buscar el archivo cuando lo que se teclea no existe ahí. El guion vive
# dentro del repositorio y la hoja suele estar al lado, en la carpeta de
# arriba: quien se para en una y nombra la otra acaba con un error que habla
# de rutas y no de lo que pasó.
DONDE_BUSCAR = (Path.cwd(), RAIZ, RAIZ.parent, Path.home() / "Documents")


def clave(texto) -> str:
    """Nombre comparable: sin tildes, sin espacios de más, en mayúsculas."""
    s = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def leer_hoja(archivo: Path) -> tuple[list[dict], Counter]:
    try:
        import openpyxl
    except ImportError:
        print(f"{ROJO}Falta openpyxl.{FIN} Instálelo con:  pip install openpyxl")
        raise SystemExit(2)

    libro = openpyxl.load_workbook(archivo, data_only=True)

    # 1. El nombre exacto, ignorando mayúsculas y tildes. Es el caso normal y
    #    resuelve la ambigüedad con la otra pestaña parecida del mismo libro.
    exactas = [h for h in libro.sheetnames if clave(h) == clave(HOJA)]
    if len(exactas) == 1:
        elegida = exactas[0]
    else:
        # 2. Si el nombre cambió —alguien corrige la errata y le pone la ese—,
        #    se acepta el parecido, pero solo si no hay dos: entre dos hojas
        #    que podrían ser, elegir es adivinar, y adivinar mal aquí no se
        #    nota nunca.
        parecidas = [h for h in libro.sheetnames
                     if clave(h).startswith(HOJA_PREFIJO)]
        if len(parecidas) == 1:
            elegida = parecidas[0]
            print(f"{AMARILLO}No hay una hoja «{HOJA}»; se usa «{elegida}», "
                  f"que es la única parecida.{FIN}")
        else:
            if not parecidas:
                print(f"{ROJO}El archivo no tiene la hoja «{HOJA}».{FIN}")
            else:
                print(f"{ROJO}Hay {len(parecidas)} hojas que podrían ser y "
                      f"ninguna se llama «{HOJA}»:{FIN}")
                for h in parecidas:
                    print(f"      «{h}»")
                print(f"   Renombre la buena como «{HOJA}», o quite las que "
                      "sobren. Cargar la equivocada no da error: da datos de otra.")
            print(f"   El archivo tiene: {', '.join(libro.sheetnames)}")
            raise SystemExit(2)

    hoja = libro[elegida]
    print(f"{GRIS}Archivo: {archivo}{FIN}")
    print(f"{GRIS}Hoja:    «{elegida}»{FIN}")

    # Que las columnas sigan donde estaban. Cargar por posición sin comprobar
    # el encabezado es la forma más silenciosa de meter fechas en el campo de
    # los días.
    encabezado = [c.value for c in hoja[1]]
    for i, esperado in ENCABEZADOS.items():
        real = str(encabezado[i] or "").strip()
        if clave(real) != clave(esperado):
            print(f"{ROJO}La columna {i + 1} dice «{real}» y se esperaba «{esperado}».{FIN}")
            print("   La hoja cambió de forma: revise el archivo antes de cargar nada.")
            raise SystemExit(2)

    eventos, descartes = [], Counter()
    for n, f in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        nombre = f[COL["nombre"]]
        if not nombre:
            continue
        ini, fin, dias = f[COL["inicio"]], f[COL["fin"]], f[COL["dias"]]
        if ini is None and fin is None and dias is None:
            continue                       # fila de persona sin vacación anotada

        if not isinstance(ini, dt.datetime) or not isinstance(fin, dt.datetime):
            descartes["fecha ausente o no válida"] += 1
            continue
        if fin < ini:
            descartes["fecha de fin anterior al inicio"] += 1
            continue
        try:
            d = float(dias)
        except (TypeError, ValueError):
            descartes["días gozados no numérico"] += 1
            continue
        if d <= 0:
            descartes["días gozados cero o negativo"] += 1
            continue

        eventos.append({
            "fila": n, "nombre": nombre, "clave": clave(nombre),
            "estado": (f[COL["estado"]] or "").strip() or None,
            "inicio": ini.date(), "fin": fin.date(), "dias": d,
        })
    return eventos, descartes


async def principal(archivo: Path, aplicar: bool, informe: Path | None) -> int:
    from app.config import get_settings
    from app.db import cerrar_pool, ejecutar, obtener_todos, obtener_uno

    if get_settings().es_produccion and not aplicar:
        pass                                # el ensayo se permite siempre

    print(f"Leyendo {archivo.name}…")
    eventos, descartes = leer_hoja(archivo)
    print(f"  {len(eventos)} vacaciones anotadas en la hoja")
    for motivo, n in descartes.most_common():
        print(f"  {AMARILLO}{n} fila(s) descartadas: {motivo}{FIN}")

    # ---------------------------------------------------------- emparejar
    gente = await obtener_todos(
        "select id, cedula, nombre from public.users where activo")
    porNombre: dict[str, list[dict]] = {}
    for p in gente:
        porNombre.setdefault(clave(p["nombre"]), []).append(p)

    repetidos = {k for k, v in porNombre.items() if len(v) > 1}
    if repetidos:
        print(f"\n{AMARILLO}Hay nombres repetidos en el sistema. Sus vacaciones NO se "
              f"cargan, porque no hay forma de saber a cuál de los dos corresponden:{FIN}")
        for k in sorted(repetidos):
            for p in porNombre[k]:
                print(f"    {p['cedula']}  {p['nombre']}")

    asignados, huerfanos = [], {}
    for e in eventos:
        candidatos = porNombre.get(e["clave"], [])
        if len(candidatos) == 1:
            asignados.append({**e, "user_id": candidatos[0]["id"],
                              "cedula": candidatos[0]["cedula"]})
        else:
            h = huerfanos.setdefault(e["clave"], {"nombre": e["nombre"],
                                                  "estado": e["estado"], "eventos": 0,
                                                  "ambiguo": len(candidatos) > 1})
            h["eventos"] += 1

    personas = {a["user_id"] for a in asignados}
    print(f"\n== EMPAREJAMIENTO ==")
    print(f"  {VERDE}{len(asignados)}{FIN} vacaciones de {VERDE}{len(personas)}{FIN} "
          f"persona(s) con nombre idéntico en el sistema")
    if huerfanos:
        total = sum(h["eventos"] for h in huerfanos.values())
        print(f"  {AMARILLO}{total}{FIN} vacaciones de {AMARILLO}{len(huerfanos)}{FIN} "
              f"nombre(s) que no empareja(n). No se cargan:")
        for k, h in sorted(huerfanos.items(), key=lambda x: -x[1]["eventos"]):
            razon = "nombre repetido en el sistema" if h["ambiguo"] else "no está en el sistema"
            print(f"     {h['eventos']:>3}  {h['estado'] or '?':<7} {h['nombre'][:40]:<40} {GRIS}{razon}{FIN}")

    # ----------------------------------------------------------- qué falta
    ya = await obtener_todos(
        """select user_id::text as user_id, fecha_inicio, fecha_fin
             from public.vacaciones_historicas""")
    existentes = {(f["user_id"], f["fecha_inicio"], f["fecha_fin"]) for f in ya}
    nuevos = [a for a in asignados
              if (str(a["user_id"]), a["inicio"], a["fin"]) not in existentes]

    # Filas de la hoja que repiten persona y fechas. Se cargan una sola vez,
    # pero se listan: o alguien tecleó dos veces lo mismo, o son dos períodos
    # distintos anotados con las mismas fechas. Ninguna de las dos cosas la
    # resuelve un programa.
    vistas, repetidas = set(), []
    for a in asignados:
        llave = (a["user_id"], a["inicio"], a["fin"])
        (repetidas.append(a) if llave in vistas else vistas.add(llave))

    print(f"\n== QUÉ SE CARGARÍA ==")
    print(f"  ya estaban cargadas: {len(asignados) - len(nuevos)}")
    if repetidas:
        print(f"  {AMARILLO}{len(repetidas)} fila(s) repiten persona y fechas en la hoja. "
              f"Se carga una sola vez; revíselas:{FIN}")
        for a in repetidas:
            print(f"     fila {a['fila']:>5}  {a['nombre'][:34]:<34} "
                  f"{a['inicio']} a {a['fin']}  ({a['dias']:g} día(s))")
    print(f"  {VERDE}nuevas: {len(nuevos)}{FIN}")

    if informe:
        with informe.open("w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["fila_hoja", "cedula", "nombre_en_la_hoja", "inicio", "fin", "dias"])
            for a in sorted(asignados, key=lambda x: (x["nombre"], x["inicio"])):
                w.writerow([a["fila"], a["cedula"], a["nombre"],
                            a["inicio"], a["fin"], f"{a['dias']:g}"])
        print(f"\n  Detalle de lo emparejado en {informe}")

    if not aplicar:
        print(f"\n{AMARILLO}Nada se escribió.{FIN} Repita con --aplicar para cargarlo.")
        return 0

    # ------------------------------------------------------------- cargar
    saldo_antes = await obtener_uno(
        "select coalesce(sum(dias_vacaciones), 0) as total from public.users")

    cargados = 0
    for a in nuevos:
        try:
            await ejecutar(
                """insert into public.vacaciones_historicas
                     (user_id, fecha_inicio, fecha_fin, dias, fila_origen)
                   values (%s, %s, %s, %s, %s)
                   on conflict (user_id, fecha_inicio, fecha_fin) do nothing""",
                (a["user_id"], a["inicio"], a["fin"], a["dias"], a["fila"]),
            )
            cargados += 1
        except Exception as exc:  # noqa: BLE001
            print(f"  {ROJO}fila {a['fila']} ({a['nombre']}): {str(exc)[:110]}{FIN}")

    saldo_despues = await obtener_uno(
        "select coalesce(sum(dias_vacaciones), 0) as total from public.users")

    print(f"\n{VERDE}Cargadas {cargados} vacaciones.{FIN}")
    if saldo_antes["total"] != saldo_despues["total"]:
        print(f"{ROJO}ATENCIÓN: el saldo total cambió de {saldo_antes['total']} a "
              f"{saldo_despues['total']}. No debía cambiar: el historial no toca el "
              f"saldo, que ya venía cargado de esta misma hoja.{FIN}")
        return 1
    print(f"  Saldo total sin cambios: {saldo_despues['total']} días, como debe ser.")

    # ------------------------------------------------------------- cotejo
    # Lo cargado tiene que sumar, persona por persona, lo mismo que suma la
    # columna de días gozados de la hoja. Si no cuadra, algo se atribuyó mal
    # y hay que saberlo ahora, no dentro de seis meses.
    porPersona: dict[str, float] = {}
    for a in asignados:
        llave = (str(a["user_id"]), a["inicio"], a["fin"])
        if llave in porPersona.get("__vistas__", set()):
            continue
        porPersona.setdefault("__vistas__", set()).add(llave)
        porPersona[str(a["user_id"])] = porPersona.get(str(a["user_id"]), 0) + a["dias"]
    porPersona.pop("__vistas__", None)

    guardado = await obtener_todos(
        """select user_id::text as user_id, sum(dias) as dias
             from public.vacaciones_historicas group by user_id""")
    enBase = {g["user_id"]: float(g["dias"]) for g in guardado}

    descuadres = [(u, d, enBase.get(u, 0)) for u, d in porPersona.items()
                  if abs(d - enBase.get(u, 0)) > 0.001]
    if descuadres:
        print(f"\n{ROJO}{len(descuadres)} persona(s) no cuadran con la hoja:{FIN}")
        for u, hoja_d, base_d in descuadres[:15]:
            persona = await obtener_uno(
                "select cedula, nombre from public.users where id = %s", (u,))
            print(f"   {persona['cedula']}  {persona['nombre'][:34]:<34} "
                  f"hoja {hoja_d:g} · sistema {base_d:g}")
        return 1
    print(f"  {VERDE}Cotejo: los días cargados cuadran con la hoja en las "
          f"{len(porPersona)} personas.{FIN}")
    return 0


def localizar(pedido: Path) -> Path:
    """El archivo que se pidió, buscándolo donde suele estar.

    El guion vive dentro del repositorio y la hoja de Talento Humano suele
    estar en la carpeta de arriba, al lado del repositorio. Quien se para en
    una y nombra la otra recibía «No existe», que es cierto y no ayuda: la
    hoja está, a un paso de donde se buscó.
    """
    if pedido.exists():
        return pedido

    # Sin repetir: al ejecutarlo desde la raíz del repositorio, el directorio
    # actual y la raíz son el mismo, y listarlo dos veces hace dudar de si el
    # guion sabe dónde está mirando.
    carpetas = list(dict.fromkeys(c.resolve() for c in DONDE_BUSCAR))

    for carpeta in carpetas:
        candidato = carpeta / pedido.name
        if candidato.exists():
            print(f"{AMARILLO}«{pedido}» no está ahí; se usa "
                  f"{candidato}{FIN}")
            return candidato

    print(f"{ROJO}No se encuentra «{pedido}».{FIN}")
    print(f"   Se buscó en: {', '.join(str(c) for c in carpetas)}")

    # Lo que sí hay cerca, para no dejar a nadie adivinando el nombre exacto.
    parecidos = sorted({
        str(f) for carpeta in carpetas if carpeta.is_dir()
        for f in carpeta.glob("*.xls*")
        if "VACACION" in clave(f.name)
    })
    if parecidos:
        print(f"\n{GRIS}Hojas de vacaciones que sí están cerca:{FIN}")
        for f in parecidos[:8]:
            print(f"   {f}")
    raise SystemExit(2)


async def _con_cierre(archivo, aplicar, informe) -> int:
    from app.db import cerrar_pool
    try:
        return await principal(archivo, aplicar, informe)
    finally:
        await cerrar_pool()


if __name__ == "__main__":
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("archivo", type=Path, help="La hoja de Talento Humano (.xlsx o .xlsm)")
    p.add_argument("--aplicar", action="store_true", help="Cargar de verdad.")
    p.add_argument("--informe", type=Path, default=None,
                   help="CSV con el detalle de lo emparejado, para cotejar.")
    a = p.parse_args()
    archivo = localizar(a.archivo)
    raise SystemExit(asyncio.run(_con_cierre(archivo, a.aplicar, a.informe)))
